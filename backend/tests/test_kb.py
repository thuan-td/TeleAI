import io
import shutil

import pytest
from docx import Document as DocxDocument
from fpdf import FPDF
from PIL import Image, ImageDraw

TESSERACT_AVAILABLE = shutil.which("tesseract") is not None


def _make_pdf_bytes(text: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, text)
    return bytes(pdf.output())


def _make_docx_bytes(text: str) -> bytes:
    doc = DocxDocument()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_image_bytes(text: str) -> bytes:
    image = Image.new("RGB", (400, 100), color="white")
    draw = ImageDraw.Draw(image)
    draw.text((10, 40), text, fill="black")
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    return buf.getvalue()


def test_create_document_then_list_returns_it(client, db_session):
    # Act
    create_response = client.post(
        "/kb/documents", json={"title": "Refund Policy", "content": "Refunds within 30 days."}
    )
    list_response = client.get("/kb/documents")

    # Assert
    assert create_response.status_code == 201
    doc_id = create_response.json()["id"]
    assert list_response.status_code == 200
    titles = [doc["title"] for doc in list_response.json()]
    assert "Refund Policy" in titles
    assert any(doc["id"] == doc_id for doc in list_response.json())


def test_delete_document_removes_it_from_list(client, db_session):
    # Arrange
    created = client.post("/kb/documents", json={"title": "Temp", "content": "Delete me"}).json()

    # Act
    delete_response = client.delete(f"/kb/documents/{created['id']}")
    list_response = client.get("/kb/documents")

    # Assert
    assert delete_response.status_code == 204
    assert all(doc["id"] != created["id"] for doc in list_response.json())


def test_delete_nonexistent_document_returns_404(client, db_session):
    # Act
    response = client.delete("/kb/documents/00000000-0000-0000-0000-000000000000")

    # Assert
    assert response.status_code == 404


def test_query_documents_matches_by_title_and_content(client, db_session):
    # Arrange
    client.post("/kb/documents", json={"title": "Refund Policy", "content": "Refunds within 30 days."})
    client.post("/kb/documents", json={"title": "Shipping", "content": "Ships in 2 business days."})

    # Act
    response = client.post("/kb/query", json={"query": "refund"})

    # Assert
    assert response.status_code == 200
    results = response.json()["results"]
    assert len(results) == 1
    assert results[0]["title"] == "Refund Policy"


def test_query_documents_no_match_returns_empty_results(client, db_session):
    # Arrange
    client.post("/kb/documents", json={"title": "Shipping", "content": "Ships in 2 business days."})

    # Act
    response = client.post("/kb/query", json={"query": "nonexistent-topic-xyz"})

    # Assert
    assert response.status_code == 200
    assert response.json()["results"] == []


RETELL_WEBHOOK_HEADERS = {"X-KB-Webhook-Secret": "test-kb-webhook-secret"}


def test_retell_function_call_extracts_query_from_args(client, db_session):
    # Arrange — Retell's actual webhook payload shape: {"name", "args": {...}, "call": {...}}
    client.post("/kb/documents", json={"title": "Refund Policy", "content": "Refunds within 30 days."})

    # Act
    response = client.post(
        "/kb/retell-function-call",
        json={
            "name": "query_knowledge_base",
            "args": {"query": "refund"},
            "call": {"call_id": "call-1"},
        },
        headers=RETELL_WEBHOOK_HEADERS,
    )

    # Assert
    assert response.status_code == 200
    results = response.json()["results"]
    assert len(results) == 1
    assert results[0]["title"] == "Refund Policy"


def test_retell_function_call_missing_query_returns_422(client, db_session):
    # Act
    response = client.post(
        "/kb/retell-function-call",
        json={"name": "query_knowledge_base", "args": {}, "call": {}},
        headers=RETELL_WEBHOOK_HEADERS,
    )

    # Assert
    assert response.status_code == 422


def test_retell_function_call_without_secret_header_returns_401(client, db_session):
    # Arrange — regression: webhook is public internet-reachable, must reject
    # unauthenticated calls instead of leaking KB content (code review finding).
    client.post("/kb/documents", json={"title": "Refund Policy", "content": "Refunds within 30 days."})

    # Act
    response = client.post(
        "/kb/retell-function-call",
        json={"name": "query_knowledge_base", "args": {"query": "refund"}, "call": {}},
    )

    # Assert
    assert response.status_code == 401


def test_retell_function_call_with_wrong_secret_returns_401(client, db_session):
    # Act
    response = client.post(
        "/kb/retell-function-call",
        json={"name": "query_knowledge_base", "args": {"query": "refund"}, "call": {}},
        headers={"X-KB-Webhook-Secret": "wrong-secret"},
    )

    # Assert
    assert response.status_code == 401


def test_upload_pdf_extracts_text_and_creates_document(client, db_session):
    # Arrange — real PDF generated with fpdf2, extracted with pypdf (no mocks)
    pdf_bytes = _make_pdf_bytes("Refund policy: 30 days from purchase.")

    # Act
    response = client.post(
        "/kb/documents/upload",
        files=[("files", ("refund-policy.pdf", pdf_bytes, "application/pdf"))],
    )

    # Assert
    assert response.status_code == 201
    results = response.json()["results"]
    assert len(results) == 1
    assert results[0]["success"] is True
    assert results[0]["document"]["title"] == "refund-policy"
    assert "Refund policy: 30 days" in results[0]["document"]["content"]


def test_upload_docx_extracts_text_and_creates_document(client, db_session):
    # Arrange — real DOCX generated with python-docx (no mocks)
    docx_bytes = _make_docx_bytes("Shipping takes 2 business days.")

    # Act
    response = client.post(
        "/kb/documents/upload",
        files=[
            (
                "files",
                (
                    "shipping.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            )
        ],
    )

    # Assert
    assert response.status_code == 201
    results = response.json()["results"]
    assert len(results) == 1
    assert results[0]["success"] is True
    assert "Shipping takes 2 business days." in results[0]["document"]["content"]


def test_upload_multiple_files_processes_each_independently(client, db_session):
    # Arrange — 1 good PDF + 1 unsupported extension, batch must not fail wholesale
    pdf_bytes = _make_pdf_bytes("Good file content.")

    # Act
    response = client.post(
        "/kb/documents/upload",
        files=[
            ("files", ("good.pdf", pdf_bytes, "application/pdf")),
            ("files", ("bad.txt", b"plain text", "text/plain")),
        ],
    )

    # Assert
    assert response.status_code == 201
    results = response.json()["results"]
    assert len(results) == 2
    good = next(r for r in results if r["filename"] == "good.pdf")
    bad = next(r for r in results if r["filename"] == "bad.txt")
    assert good["success"] is True
    assert bad["success"] is False
    assert bad["error"] is not None


def test_upload_rejects_file_over_size_limit(client, db_session):
    # Arrange
    oversized = b"x" * (5 * 1024 * 1024 + 1)

    # Act
    response = client.post(
        "/kb/documents/upload",
        files=[("files", ("huge.pdf", oversized, "application/pdf"))],
    )

    # Assert
    assert response.status_code == 201  # batch endpoint always 201, per-file result carries the error
    result = response.json()["results"][0]
    assert result["success"] is False
    assert "5MB" in result["error"]


@pytest.mark.skipif(not TESSERACT_AVAILABLE, reason="tesseract binary not installed on this host")
def test_upload_image_runs_ocr_and_creates_document(client, db_session):
    # Arrange — real PNG with rendered text, real Tesseract OCR (no mocks).
    # Skipped on hosts without the tesseract binary (present in Docker image,
    # see backend/Dockerfile) — install via `brew install tesseract
    # tesseract-lang` for local dev to run this test.
    image_bytes = _make_image_bytes("HELLO WORLD")

    # Act
    response = client.post(
        "/kb/documents/upload",
        files=[("files", ("sign.png", image_bytes, "image/png"))],
    )

    # Assert
    assert response.status_code == 201
    result = response.json()["results"][0]
    assert result["success"] is True
    assert "HELLO" in result["document"]["content"].upper()
